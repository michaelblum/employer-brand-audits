from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))

from scripts.publication_pipeline_fixture import (
    SEGMENT_TVP_TEMPLATE_ID,
    generate_segment_tvp_audit_fixture,
)
from scripts.workbench_projection import project_audit_manifest


ADT_DOCX = REPO_ROOT / "reference_publications" / "Reference" / "ADT Tech Product TVP 2025 Client Data Immersion and Comp Audit.docx"
ADT_PPTX = REPO_ROOT / "reference_publications" / "Reference" / "ADT Comp Audit for Tech.pptx"


class SegmentTvpPipelineTests(unittest.TestCase):
    def arbitrary_profile(self) -> dict[str, object]:
        return {
            "project_id": "nova-systems-segment-tvp",
            "client_name": "Nova Systems",
            "client_full_name": "Nova Systems",
            "industry": "Industrial software",
            "geography": "United States",
            "audience": "Engineering and product talent",
            "domain": "novasystems.example",
            "target_talent_segments": ["Engineering", "Product"],
            "segment_competitors": [
                {"entity_id": "competitor-orbit", "name": "Orbit Labs", "careers_url": "https://orbit.example/careers"},
                {"entity_id": "competitor-pulse", "name": "Pulse Works", "careers_url": "https://pulse.example/careers"},
            ],
            "social_sources": ["https://social.example/nova"],
        }

    @unittest.skipUnless(ADT_DOCX.exists() and ADT_PPTX.exists(), "local ADT reference files are not tracked")
    def test_local_adt_references_are_readable_when_present(self) -> None:
        try:
            from docx import Document
            from pptx import Presentation
        except ImportError as exc:
            self.skipTest(f"reference readers unavailable: {exc}")

        document = Document(ADT_DOCX)
        deck = Presentation(ADT_PPTX)
        with ZipFile(ADT_DOCX) as archive:
            media_files = [name for name in archive.namelist() if name.startswith("word/media/")]

        self.assertEqual(len(document.tables), 11)
        self.assertGreaterEqual(len(media_files), 28)
        self.assertEqual(len(deck.slides), 16)

    def test_segment_tvp_fixture_writes_manifest_and_required_views(self) -> None:
        with tempfile.TemporaryDirectory(dir=REPO_ROOT / "artifacts") as tmp:
            manifest_path = generate_segment_tvp_audit_fixture(Path(tmp) / "segment-tvp")
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))

            required_files = [
                "pipeline-intake.json",
                "project-frame.json",
                "segment-source-roster.json",
                "segment-capture-pack.json",
                "segment-evidence-matrix.json",
                "tvp-analysis-pack.json",
                "report-docx-view.html",
                "brand-strategy-deck-view.html",
                "social-platform-audit-view.html",
                "data-workbook-view.html",
                "l4-publication.html",
            ]
            for relative_path in required_files:
                self.assertTrue((manifest_path.parent / relative_path).exists(), relative_path)

        self.assertEqual(manifest["template_id"], SEGMENT_TVP_TEMPLATE_ID)
        self.assertNotIn("reference_publications", json.dumps(manifest))
        self.assertEqual(
            [step["id"] for step in manifest["steps"]],
            [
                "p0-pipeline-intake",
                "p0-project-frame",
                "p0-segment-source-roster",
                "p1-segment-capture-pack",
                "p2-segment-evidence-matrix",
                "p3-tvp-analysis-pack",
                "p4-report-docx-view",
                "p4-brand-strategy-deck-view",
                "p4-social-platform-audit-view",
                "p4-data-workbook-view",
                "p4-l4-publication-view",
            ],
        )

    def test_segment_tvp_intake_and_data_groups(self) -> None:
        with tempfile.TemporaryDirectory(dir=REPO_ROOT / "artifacts") as tmp:
            manifest_path = generate_segment_tvp_audit_fixture(Path(tmp) / "segment-tvp")
            root = manifest_path.parent
            intake = json.loads((root / "pipeline-intake.json").read_text(encoding="utf-8"))
            roster = json.loads((root / "segment-source-roster.json").read_text(encoding="utf-8"))
            capture_pack = json.loads((root / "segment-capture-pack.json").read_text(encoding="utf-8"))
            evidence_matrix = json.loads((root / "segment-evidence-matrix.json").read_text(encoding="utf-8"))
            analysis_pack = json.loads((root / "tvp-analysis-pack.json").read_text(encoding="utf-8"))

        self.assertEqual(intake["pipeline_id"], SEGMENT_TVP_TEMPLATE_ID)
        self.assertIn("Technology", intake["target_talent_segments"])
        self.assertIn("role_evidence_sources", intake)
        self.assertIn("segment_competitors", intake)
        self.assertIn("social_sources", intake)
        self.assertEqual(len(roster["entities"]), 9)
        self.assertGreaterEqual(len(capture_pack["source_artifacts"]), 18)
        self.assertIn("job_posting", {artifact["source_type"] for artifact in capture_pack["source_artifacts"]})
        self.assertIn("role_family_careers_page", {artifact["source_type"] for artifact in capture_pack["source_artifacts"]})
        self.assertIn("social_profile", {artifact["source_type"] for artifact in capture_pack["source_artifacts"]})
        self.assertGreaterEqual(len(evidence_matrix["evidence_items"]), 90)
        self.assertTrue(all(item["talent_segment"] for item in evidence_matrix["evidence_items"]))
        evidence_ids = {item["evidence_id"] for item in evidence_matrix["evidence_items"]}
        for finding in analysis_pack["segment_positioning_findings"] + analysis_pack["segment_recommendations"]:
            self.assertTrue(set(finding["supporting_evidence_ids"]).issubset(evidence_ids), finding)

    def test_segment_tvp_projection_groups_publication_views(self) -> None:
        with tempfile.TemporaryDirectory(dir=REPO_ROOT / "artifacts") as tmp:
            manifest_path = generate_segment_tvp_audit_fixture(Path(tmp) / "segment-tvp")
            projection = project_audit_manifest(manifest_path)

        group_slots = {group["slot"] for group in projection["artifact_groups"]}
        self.assertIn("publication.report_docx.bundle", group_slots)
        self.assertIn("publication.brand_strategy_deck.bundle", group_slots)
        self.assertIn("publication.social_platform_audit.bundle", group_slots)
        self.assertIn("publication.data_workbook.bundle", group_slots)
        self.assertIn("publication.l4_publication.bundle", group_slots)

    def test_arbitrary_segment_tvp_profile_does_not_leak_reference_labels(self) -> None:
        with tempfile.TemporaryDirectory(dir=REPO_ROOT / "artifacts") as tmp:
            root = Path(tmp)
            profile_path = root / "nova-profile.json"
            profile_path.write_text(json.dumps(self.arbitrary_profile()), encoding="utf-8")
            manifest_path = generate_segment_tvp_audit_fixture(
                root / "segment-tvp",
                project_profile_path=profile_path,
            )

            generated_text = "\n".join(
                path.read_text(encoding="utf-8")
                for path in sorted(manifest_path.parent.glob("*"))
                if path.is_file() and path.suffix in {".json", ".html"}
            )

        self.assertIn("Nova Systems", generated_text)
        banned = [
            "ADT",
            "jobs.adt.com",
            "Vivint",
            "SimpliSafe",
            "Arlo",
            "Verisure",
            "Alarm.com",
            "Resideo",
            "Qualcomm",
            "State Farm",
            "Google",
            "Uber",
            "ADT-style",
            "HarbourVest-style",
        ]
        for label in banned:
            self.assertNotIn(label, generated_text)

    def test_segment_tvp_fixture_is_registered_in_command_surface(self) -> None:
        from scripts.eba_cli import FIXTURE_GENERATORS, demo_recipe_lines, validation_commands

        self.assertIn("segment-tvp-audit", FIXTURE_GENERATORS)
        self.assertTrue(
            any(command[-1] == "tests/test_publication_segment_tvp.py" for command in validation_commands())
        )
        with tempfile.TemporaryDirectory(dir=REPO_ROOT / "artifacts") as tmp:
            manifest = Path(tmp) / "manifest.json"
            manifest.write_text(json.dumps({"template_id": SEGMENT_TVP_TEMPLATE_ID}), encoding="utf-8")
            lines = demo_recipe_lines(fixture=None, manifest=manifest)

        self.assertIn("Segment Source Roster", "\n".join(lines))
        self.assertIn("Social Platform Audit", "\n".join(lines))


if __name__ == "__main__":
    unittest.main()
